import os
import sys
from pathlib import Path
from types import SimpleNamespace
import io

from openpyxl import load_workbook

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))

import services.financial_automation_service as automation_module
from services.financial_import_service import FinancialImportService
from services.financial_automation_service import FinancialAutomationService


class _Column:
    def __eq__(self, other):
        return self

    def is_(self, other):
        return self

    def in_(self, other):
        return self

    def asc(self):
        return self


class _Batch:
    def __init__(self, batch_id=1, origin_type="manual_upload", source_label="Lote teste"):
        self.id = batch_id
        self.origin_type = origin_type
        self.source_label = source_label
        self.status_summary_json = {}
        self.updated_at = None
        self.documents = SimpleNamespace(filter=lambda *args, **kwargs: SimpleNamespace(count=lambda: 0))
        self.records = SimpleNamespace(filter=lambda *args, **kwargs: SimpleNamespace(all=lambda: []))

    def __hash__(self):
        return hash(self.id)

    def to_dict(self):
        return {
            "id": self.id,
            "origin_type": self.origin_type,
            "source_label": self.source_label,
            "status_summary_json": self.status_summary_json,
        }


class _Record:
    def __init__(self, record_id, settlement_state, batch=None):
        self.id = record_id
        self.company_id = 9
        self.batch_id = 1
        self.batch = batch or _Batch()
        self.status = "validated"
        self.entry_direction = "payable"
        self.settlement_state = settlement_state
        self.description = f"Registro {record_id}"
        self.counterparty_id = None
        self.bank_account_id = None
        self.chart_account_id = None
        self.cost_center_id = None
        self.domain_type = None
        self.domain_source_id = None
        self.amount = 150
        self.competence_date = None
        self.due_date = None
        self.confidence_score = 0.8
        self.validation_notes = None
        self.normalized_payload_json = {}
        self.metadata_json = {}
        self.generated_financial_entry_id = None
        self.generated_financial_schedule_id = None
        self.validated_by_user_id = None
        self.validated_at = None
        self.generated_by_user_id = None
        self.generated_at = None
        self.created_at = None
        self.updated_at = None
        self.deleted_at = None
        self.source_document = None

    def to_dict(self):
        return {
            "id": self.id,
            "company_id": self.company_id,
            "batch_id": self.batch_id,
            "status": self.status,
            "entry_direction": self.entry_direction,
            "settlement_state": self.settlement_state,
            "description": self.description,
            "counterparty_id": self.counterparty_id,
            "bank_account_id": self.bank_account_id,
            "chart_account_id": self.chart_account_id,
            "cost_center_id": self.cost_center_id,
            "domain_type": self.domain_type,
            "domain_source_id": self.domain_source_id,
            "amount": self.amount,
            "competence_date": self.competence_date,
            "due_date": self.due_date,
            "confidence_score": self.confidence_score,
            "validation_notes": self.validation_notes,
            "normalized_payload_json": self.normalized_payload_json,
            "metadata_json": self.metadata_json,
            "generated_financial_entry_id": self.generated_financial_entry_id,
            "generated_financial_schedule_id": self.generated_financial_schedule_id,
        }


def test_automation_center_template_contains_bulk_validation_flow():
    template = Path(
        r"C:\GestaoVersus\app32\app32\templates\modules\financial\automation_center.html"
    ).read_text(encoding="utf-8")

    assert "Central de Automação Financeira" in template
    assert 'id="fa-mark-validated"' in template
    assert 'id="fa-mark-excluded"' in template
    assert 'id="fa-generate"' in template
    assert "Importada" in template
    assert "Validada" in template
    assert "Gerada" in template
    assert "Excluída" in template
    assert 'id="fa-import-files"' in template
    assert 'id="filter-batch"' in template
    assert 'id="filter-document-type"' in template
    assert "<th>Lote</th>" in template
    assert "<th>Documento</th>" not in template
    assert "XML fiscal" in template
    assert "Baixar planilha modelo" in template
    assert "/financial/automation/template" in template


def test_automation_center_js_uses_data_field_mapping_and_origin_labels():
    script = Path(
        r"C:\GestaoVersus\app32\app32\static\js\financial_automation_center.js"
    ).read_text(encoding="utf-8")

    assert "statusLabels" in script
    assert "originLabels" in script
    assert "row.querySelectorAll('[data-field]')" in script
    assert "'domain_link'" in script
    assert "record.batch?.source_label" in script
    assert "batch_options" in script
    assert "batchDisplayCode" in script
    assert "filter-batch" in script
    assert "batch.source_label ? ` ·" not in script
    assert "/api/financial/automation/uploads" in script
    assert "new FormData()" in script
    assert "documentTypeLabels" in script
    assert "filter-document-type" in script
    assert "related_documents" in script
    assert "method: 'DELETE'" in script
    assert "/api/financial/automation/records/${id}?company_id=${companyId}" in script


def test_delete_record_soft_deletes_when_no_active_links(monkeypatch):
    batch = _Batch()
    item = _Record(record_id=41, settlement_state="open", batch=batch)

    class _Query:
        def filter(self, *args, **kwargs):
            return self

        def first(self):
            return item

    class _RecordModel:
        id = _Column()
        company_id = _Column()
        deleted_at = _Column()
        query = _Query()

    monkeypatch.setattr(automation_module, "FinancialAutomationRecord", _RecordModel)
    monkeypatch.setattr(automation_module.FinancialService, "_ensure_company_scope", lambda *args, **kwargs: None)
    monkeypatch.setattr(automation_module.FinancialAutomationService, "_append_history", lambda **kwargs: None)
    monkeypatch.setattr(automation_module.FinancialAutomationService, "_refresh_batch_summary", lambda batch: None)
    monkeypatch.setattr(automation_module.db.session, "commit", lambda: None)
    monkeypatch.setattr(automation_module.db.session, "rollback", lambda: None)

    result, error = FinancialAutomationService.delete_record(
        company_id=9,
        record_id=41,
        allowed_company_ids=[9],
        performed_by_user_id=77,
    )

    assert error is None
    assert result == {"id": 41, "deleted": True}
    assert item.status == "excluded"
    assert item.deleted_at is not None


def test_delete_record_blocks_when_generated_links_are_active(monkeypatch):
    batch = _Batch()
    item = _Record(record_id=42, settlement_state="open", batch=batch)
    item.generated_financial_entry_id = 501

    class _Query:
        def filter(self, *args, **kwargs):
            return self

        def first(self):
            return item

    class _RecordModel:
        id = _Column()
        company_id = _Column()
        deleted_at = _Column()
        query = _Query()

    monkeypatch.setattr(automation_module, "FinancialAutomationRecord", _RecordModel)
    monkeypatch.setattr(automation_module.FinancialService, "_ensure_company_scope", lambda *args, **kwargs: None)
    monkeypatch.setattr(
        automation_module.FinancialAutomationService,
        "_list_delete_blockers",
        lambda current: [{"type": "financial_entry", "id": 501, "label": "Lançamento financeiro #501"}],
    )

    result, error = FinancialAutomationService.delete_record(
        company_id=9,
        record_id=42,
        allowed_company_ids=[9],
        performed_by_user_id=77,
    )

    assert result is None
    assert "vínculos financeiros ativos" in error
    assert "Lançamento financeiro #501" in error


def test_financial_import_service_accepts_xls_dispatch(monkeypatch):
    import services.financial_import_service as import_module

    monkeypatch.setattr(
        import_module.FinancialImportService,
        "_parse_xls_bytes",
        lambda file_bytes: [{"descricao": "Linha XLS"}],
    )

    rows = import_module.FinancialImportService._parse_source_rows("xls", b"fake-xls")

    assert rows == [{"descricao": "Linha XLS"}]


def test_financial_import_template_uses_instruction_layout_with_horizontal_logo():
    content, error = FinancialImportService.build_import_template()

    assert error is None
    workbook = load_workbook(io.BytesIO(content))
    sheet = workbook["Instruções"]
    data_sheet = workbook["Importação"]

    assert "A1:D1" in {str(item) for item in sheet.merged_cells.ranges}
    assert "A2:D2" in {str(item) for item in sheet.merged_cells.ranges}
    assert "A5:I5" in {str(item) for item in sheet.merged_cells.ranges}
    assert "F1:I2" in {str(item) for item in sheet.merged_cells.ranges}
    assert "B8:I8" in {str(item) for item in sheet.merged_cells.ranges}
    assert sheet["A1"].value == "Versus Gestão Corporativa"
    assert sheet["A2"].value == "Modelo de importação financeira APP32"
    assert sheet["F1"].fill.fgColor.rgb in {"00000000", "000000"}
    assert sheet["A8"].value == "1."
    assert "conta a pagar / receber" in sheet["B11"].value
    assert "Pagar" in sheet["B12"].value
    assert "código completo ou código reduzido" in sheet["B13"].value
    assert sheet["A22"].value == "Exemplo de preenchimento"
    assert sheet["A23"].value == "Tipo de Registro *"
    assert sheet["B23"].value == "Tipo do Título *"
    assert sheet["A24"].value == "agendamento"
    assert sheet["B24"].value == "Pagar ou Receber"
    assert data_sheet["A1"].value == "Tipo de Registro *"
    assert data_sheet["B1"].value == "Tipo do Título *"
    assert data_sheet["A2"].value is None
    assert data_sheet["B2"].value is None
    assert data_sheet.freeze_panes == "A2"


def test_generate_records_routes_settled_to_entry_and_open_to_schedule(monkeypatch):
    batch = _Batch()
    settled_record = _Record(record_id=1, settlement_state="settled", batch=batch)
    open_record = _Record(record_id=2, settlement_state="open", batch=batch)
    records = [settled_record, open_record]

    class _Query:
        def filter(self, *args, **kwargs):
            return self

        def order_by(self, *args, **kwargs):
            return self

        def all(self):
            return records

    class _RecordModel:
        company_id = _Column()
        deleted_at = _Column()
        status = _Column()
        id = _Column()
        query = _Query()

    monkeypatch.setattr(automation_module, "FinancialAutomationRecord", _RecordModel)
    monkeypatch.setattr(
        automation_module.FinancialAutomationGenerateInput,
        "model_validate",
        lambda payload: SimpleNamespace(company_id=9, record_ids=None, only_status="validated", generated_by_user_id=77),
    )
    monkeypatch.setattr(automation_module.FinancialService, "_ensure_company_scope", lambda *args, **kwargs: None)
    monkeypatch.setattr(automation_module.FinancialAutomationService, "_generate_entry_from_record", lambda *args, **kwargs: (501, None))
    monkeypatch.setattr(automation_module.FinancialAutomationService, "_generate_schedule_from_record", lambda *args, **kwargs: (901, None))
    monkeypatch.setattr(automation_module.FinancialAutomationService, "_append_history", lambda **kwargs: None)
    monkeypatch.setattr(automation_module.FinancialAutomationService, "_refresh_batch_summary", lambda batch: None)
    monkeypatch.setattr(automation_module.db.session, "commit", lambda: None)
    monkeypatch.setattr(automation_module.db.session, "rollback", lambda: None)

    result, error = FinancialAutomationService.generate_records(
        payload={"company_id": 9},
        allowed_company_ids=[9],
    )

    assert error is None
    assert result["count"] == 2
    assert settled_record.generated_financial_entry_id == 501
    assert open_record.generated_financial_schedule_id == 901
    assert settled_record.status == "generated"
    assert open_record.status == "generated"


def test_bulk_update_status_rejects_manual_generated_transition(monkeypatch):
    batch = _Batch()
    item = _Record(record_id=1, settlement_state="open", batch=batch)

    class _Query:
        def filter(self, *args, **kwargs):
            return self

        def all(self):
            return [item]

    class _RecordModel:
        company_id = _Column()
        id = _Column()
        deleted_at = _Column()
        query = _Query()

    monkeypatch.setattr(automation_module, "FinancialAutomationRecord", _RecordModel)
    monkeypatch.setattr(
        automation_module.FinancialAutomationBulkStatusInput,
        "model_validate",
        lambda payload: SimpleNamespace(company_id=9, record_ids=[1], status="generated", validation_notes=None),
    )
    monkeypatch.setattr(automation_module.FinancialService, "_ensure_company_scope", lambda *args, **kwargs: None)

    result, error = FinancialAutomationService.bulk_update_status(
        payload={"company_id": 9, "record_ids": [1], "status": "generated"},
        allowed_company_ids=[9],
    )

    assert result is None
    assert error == "Use a ação de geração oficial para marcar registros como Gerada."


def test_upload_batch_files_creates_documents_without_records(monkeypatch):
    added = []

    class _Session:
        def add(self, obj):
            added.append(obj)

        def flush(self):
            for index, obj in enumerate(added, start=1):
                if getattr(obj, "id", None) is None:
                    obj.id = index

        def commit(self):
            return None

        def rollback(self):
            return None

    class _BatchModel:
        def __init__(self, **kwargs):
            self.__dict__.update(kwargs)
            self.id = None
            self.status_summary_json = kwargs.get("status_summary_json", {})

        def to_dict(self):
            return {
                "id": self.id,
                "company_id": self.company_id,
                "origin_type": self.origin_type,
                "source_label": self.source_label,
                "status_summary_json": self.status_summary_json,
            }

    class _DocumentModel:
        def __init__(self, **kwargs):
            self.__dict__.update(kwargs)
            self.id = None
            self.created_at = None
            self.updated_at = None
            self.deleted_at = None

        def to_dict(self):
            return {
                "id": self.id,
                "company_id": self.company_id,
                "batch_id": self.batch_id,
                "file_name": self.file_name,
                "stored_relative_path": self.stored_relative_path,
                "mime_type": self.mime_type,
                "document_type": getattr(self, "document_type", None),
                "document_group_key": getattr(self, "document_group_key", None),
            }

    monkeypatch.setattr(automation_module.db, "session", _Session())
    monkeypatch.setattr(automation_module, "FinancialAutomationBatch", _BatchModel)
    monkeypatch.setattr(automation_module, "FinancialAutomationDocument", _DocumentModel)
    monkeypatch.setattr(automation_module.FinancialService, "_ensure_company_scope", lambda *args, **kwargs: None)
    monkeypatch.setattr(
        automation_module.FinancialAccountabilityService,
        "store_document",
        lambda **kwargs: (
            {
                "file_name": "nf_teste.pdf",
                "stored_relative_path": "financial/automation/9/nf_teste.pdf",
                "public_url": "/uploads/financial/automation/9/nf_teste.pdf",
                "mime_type": "application/pdf",
                "file_size": 100,
                "file_size_original": 100,
                "file_size_optimized": 80,
                "sha256": "abc",
                "extracted_text": "nota fiscal teste",
                "extracted_preview": "nota fiscal teste",
                "extraction_method": "pdf_text",
                "extension": ".pdf",
                "document_family": "fiscal",
                "document_type": "danfe_pdf",
                "source_kind": "pdf",
                "parser_status": "parsed",
                "parser_version": "v2",
                "document_group_key": "key:123",
                "confidence_score": 0.84,
                "structured_payload_json": {"document_number": "123"},
                "original_relative_path": "financial/automation/9/original/nf_teste.pdf",
                "optimized_relative_path": "financial/automation/9/derived/nf_teste_optimized.pdf",
                "preview_relative_path": "financial/automation/9/derived/nf_teste_preview.webp",
                "original_public_url": "/uploads/financial/automation/9/original/nf_teste.pdf",
                "optimized_public_url": "/uploads/financial/automation/9/derived/nf_teste_optimized.pdf",
                "preview_public_url": "/uploads/financial/automation/9/derived/nf_teste_preview.webp",
            },
            None,
        ),
    )

    result, error = FinancialAutomationService.upload_batch_files(
        company_id=9,
        origin_type="manual_upload",
        files=[object()],
        upload_root="C:/tmp/uploads",
        source_label="Lote upload real",
        created_by_user_id=7,
        allowed_company_ids=[9],
    )

    assert error is None
    assert result is not None
    assert result["batch"]["origin_type"] == "manual_upload"
    assert result["batch"]["status_summary_json"]["documents_total"] == 1
    assert result["records"] == []
    assert len(result["documents"]) == 1
    assert result["documents"][0]["document_type"] == "danfe_pdf"
    assert result["documents"][0]["document_group_key"] == "key:123"


def test_parse_batch_documents_groups_xml_and_danfe_into_single_record(monkeypatch):
    added = []

    class _Session:
        def add(self, obj):
            added.append(obj)

        def flush(self):
            for index, obj in enumerate(added, start=1):
                if getattr(obj, "id", None) is None:
                    obj.id = 100 + index

        def commit(self):
            return None

        def rollback(self):
            return None

    class _RecordModel:
        deleted_at = _Column()

        def __init__(self, **kwargs):
            self.__dict__.update(kwargs)
            self.id = None
            self.generated_financial_entry_id = None
            self.generated_financial_schedule_id = None
            self.validated_by_user_id = None
            self.validated_at = None
            self.generated_by_user_id = None
            self.generated_at = None
            self.created_at = None
            self.updated_at = None
            self.deleted_at = None
            self.source_document = None
            self.batch = None

        def to_dict(self):
            return {
                "id": self.id,
                "company_id": self.company_id,
                "batch_id": self.batch_id,
                "status": self.status,
                "entry_direction": self.entry_direction,
                "settlement_state": self.settlement_state,
                "description": self.description,
                "amount": float(self.amount),
                "document_group_key": self.document_group_key,
                "document_type": self.document_type,
                "document_key": self.document_key,
                "external_document_number": self.external_document_number,
                "issuer_name": self.issuer_name,
                "issuer_document": self.issuer_document,
                "recipient_name": self.recipient_name,
                "recipient_document": self.recipient_document,
                "issue_date": self.issue_date.isoformat() if self.issue_date else None,
                "confidence_score": float(self.confidence_score),
                "validation_notes": self.validation_notes,
                "extracted_fields_json": self.extracted_fields_json,
                "review_flags_json": self.review_flags_json,
                "normalized_payload_json": self.normalized_payload_json,
                "metadata_json": self.metadata_json,
            }

    class _DocRecords:
        def filter(self, *args, **kwargs):
            return self

        def count(self):
            return 0

    class _BatchRecords:
        def filter(self, *args, **kwargs):
            return self

        def all(self):
            return []

    class _BatchDocs:
        def __init__(self, docs):
            self._docs = docs

        def filter(self, *args, **kwargs):
            return self

        def all(self):
            return list(self._docs)

        def count(self):
            return len(self._docs)

    class _Document:
        deleted_at = _Column()
        company_id = _Column()
        batch_id = _Column()
        document_group_key = _Column()
        id = _Column()
        query = SimpleNamespace(filter=lambda *args, **kwargs: SimpleNamespace(order_by=lambda *a, **k: SimpleNamespace(all=lambda: [])))

        def __init__(self, doc_id, document_type):
            self.id = doc_id
            self.company_id = 9
            self.batch_id = 1
            self.file_name = f"doc_{doc_id}.{'xml' if document_type.endswith('_xml') else 'pdf'}"
            self.document_type = document_type
            self.source_kind = "xml" if document_type.endswith("_xml") else "pdf"
            self.document_group_key = "key:35260412345678000199550010000012341000012345"
            self.confidence_score = 0.98 if document_type.endswith("_xml") else 0.84
            self.parser_status = "parsed"
            self.structured_payload_json = {
                "document_key": "35260412345678000199550010000012341000012345",
                "document_number": "1234",
                "document_series": "1",
                "issuer_name": "Empresa Emitente LTDA",
                "issuer_document": "12345678000199",
                "recipient_name": "Cliente Teste SA",
                "recipient_document": "99887766000155",
                "issue_date": "2026-04-17",
                "total_amount": 1500.0,
            }
            self.preview_payload_json = {}
            self.extracted_text = "DANFE" if document_type == "danfe_pdf" else "<xml/>"
            self.records = _DocRecords()
            self.deleted_at = None

        def to_dict(self):
            return {
                "id": self.id,
                "document_type": self.document_type,
                "document_group_key": self.document_group_key,
                "file_name": self.file_name,
            }

    batch = _Batch()
    xml_doc = _Document(1, "nfe_xml")
    danfe_doc = _Document(2, "danfe_pdf")
    batch.documents = _BatchDocs([xml_doc, danfe_doc])
    batch.records = _BatchRecords()

    monkeypatch.setattr(automation_module.db, "session", _Session())
    monkeypatch.setattr(automation_module, "FinancialAutomationRecord", _RecordModel)
    monkeypatch.setattr(automation_module, "FinancialAutomationDocument", _Document)
    monkeypatch.setattr(automation_module.FinancialService, "_ensure_company_scope", lambda *args, **kwargs: None)
    monkeypatch.setattr(automation_module.FinancialAutomationService, "_load_batch_for_company", lambda *args, **kwargs: batch)
    monkeypatch.setattr(automation_module.FinancialAutomationService, "_read_document_bytes", lambda *args, **kwargs: b"fake")
    monkeypatch.setattr(automation_module.FinancialAutomationService, "_append_history", lambda **kwargs: None)
    monkeypatch.setattr(automation_module.FinancialAutomationService, "_refresh_batch_summary", lambda batch: None)
    monkeypatch.setattr(
        automation_module.FinancialAutomationService,
        "_apply_auto_suggestions_to_record_kwargs",
        lambda **kwargs: kwargs["record_kwargs"],
    )
    monkeypatch.setattr(
        automation_module.FinancialAutomationService,
        "_apply_duplicate_metadata_to_record_kwargs",
        lambda **kwargs: kwargs["record_kwargs"],
    )

    result, error = FinancialAutomationService.parse_batch_documents(
        company_id=9,
        batch_id=1,
        upload_root="C:/tmp/uploads",
        allowed_company_ids=[9],
        performed_by_user_id=7,
    )

    assert error is None
    assert result["count"] == 1
    assert result["records"][0]["document_type"] == "nfe_xml"
    assert result["records"][0]["document_group_key"].startswith("key:")


def test_build_record_kwargs_from_document_group_applies_suggestions_and_dedupe(monkeypatch):
    class _Document:
        id = 10
        document_type = "danfe_pdf"
        document_group_key = "key:123"
        confidence_score = 0.88
        sha256 = "abc123"
        parser_status = "parsed"
        structured_payload_json = {
            "document_key": "123",
            "document_number": "NF-1",
            "issuer_name": "Fornecedor XPTO",
            "issuer_document": "12345678000199",
            "recipient_name": "Cliente",
            "recipient_document": "99887766000155",
            "issue_date": "2026-04-17",
            "total_amount": 99.9,
        }

    monkeypatch.setattr(
        automation_module.FinancialAutomationService,
        "_apply_auto_suggestions_to_record_kwargs",
        lambda **kwargs: {**kwargs["record_kwargs"], "counterparty_id": 77, "chart_account_id": 88, "cost_center_id": 99},
    )
    monkeypatch.setattr(
        automation_module.FinancialAutomationService,
        "_apply_duplicate_metadata_to_record_kwargs",
        lambda **kwargs: {
            **kwargs["record_kwargs"],
            "review_flags_json": list(kwargs["record_kwargs"].get("review_flags_json") or []) + ["duplicate_detected"],
            "metadata_json": {
                **dict(kwargs["record_kwargs"].get("metadata_json") or {}),
                "dedupe": {"status": "duplicate", "reason": "document_key", "matched_record_id": 501},
            },
        },
    )

    payload = FinancialAutomationService._build_record_kwargs_from_document_group(
        company_id=9,
        batch_id=3,
        anchor_document=_Document(),
        grouped_documents=[_Document()],
    )

    assert payload["counterparty_id"] == 77
    assert payload["chart_account_id"] == 88
    assert payload["cost_center_id"] == 99
    assert "duplicate_detected" in payload["review_flags_json"]
    assert payload["metadata_json"]["dedupe"]["matched_record_id"] == 501


def test_generate_records_blocks_exact_duplicates(monkeypatch):
    batch = _Batch()
    duplicated = _Record(record_id=3, settlement_state="open", batch=batch)
    duplicated.metadata_json = {"dedupe": {"status": "duplicate", "reason": "sha256"}}

    class _Query:
        def filter(self, *args, **kwargs):
            return self

        def order_by(self, *args, **kwargs):
            return self

        def all(self):
            return [duplicated]

    class _RecordModel:
        company_id = _Column()
        deleted_at = _Column()
        status = _Column()
        id = _Column()
        query = _Query()

    monkeypatch.setattr(automation_module, "FinancialAutomationRecord", _RecordModel)
    monkeypatch.setattr(
        automation_module.FinancialAutomationGenerateInput,
        "model_validate",
        lambda payload: SimpleNamespace(company_id=9, record_ids=[3], only_status="validated", generated_by_user_id=77),
    )
    monkeypatch.setattr(automation_module.FinancialService, "_ensure_company_scope", lambda *args, **kwargs: None)

    result, error = FinancialAutomationService.generate_records(
        payload={"company_id": 9, "record_ids": [3]},
        allowed_company_ids=[9],
    )

    assert result is None
    assert "duplicidade exata" in error.lower()
